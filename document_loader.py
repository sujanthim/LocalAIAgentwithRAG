import os
import uuid
import csv
import json

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from role_config import get_allowed_roles


class DocumentLoader:
    # Text splitting configuration
    CHUNK_SIZE = 1000  # Maximum characters per chunk
    CHUNK_OVERLAP = 200  # Overlap between chunks to maintain context

    @classmethod
    def load_single_file(cls, filepath: str) -> tuple[list[Document], list[str]]:
        """Load a single file and return (documents, ids). Stamps allowed_roles metadata."""
        ext = os.path.splitext(filepath)[1].lower()
        basename = os.path.basename(filepath)
        allowed_roles = get_allowed_roles(basename)

        documents = []

        if ext == ".pdf":
            documents = cls._load_pdf(filepath)
        elif ext == ".docx":
            documents = cls._load_docx(filepath)
        elif ext == ".csv":
            documents = cls._load_csv(filepath)
        elif ext in (".xlsx", ".xls"):
            documents = cls._load_excel(filepath)
        elif ext == ".json":
            documents = cls._load_json(filepath)
        elif ext in (".txt", ".md"):
            documents = cls._load_text(filepath)
        else:
            print(f"Unsupported file type: {ext} for {filepath}")
            return [], []

        for doc in documents:
            doc.metadata["source"] = filepath
            doc.metadata["file_type"] = ext
            # Store allowed_roles as JSON string (ChromaDB doesn't support lists)
            doc.metadata["allowed_roles"] = json.dumps(allowed_roles)

        # Row-based formats are already one record per Document — skip chunking
        if ext not in (".csv", ".xlsx", ".xls", ".json"):
            documents = cls._split_documents(documents)
        
        ids = [str(uuid.uuid4()) for _ in documents]
        return documents, ids

    @classmethod
    def load_documents_from_directory(
        cls,
        directory: str,
        file_types: list[str] = None,
    ) -> tuple[list[Document], list[str]]:
        """Load all matching files from a directory and return (documents, ids)."""
        if file_types is None:
            file_types = [".pdf", ".docx", ".txt", ".csv", ".md", ".xlsx", ".xls", ".json"]

        all_documents: list[Document] = []
        all_ids: list[str] = []

        for filename in sorted(os.listdir(directory)):
            filepath = os.path.join(directory, filename)
            
            # Skip directories — only process files
            if not os.path.isfile(filepath):
                continue
                
            ext = os.path.splitext(filename)[1].lower()
            if ext not in file_types:
                continue
            
            docs, ids = cls.load_single_file(filepath)
            all_documents.extend(docs)
            all_ids.extend(ids)

        ##print(f"Loaded {len(all_documents)} document chunks from {directory}")
        return all_documents, all_ids

    # ------------------------------------------------------------------
    # Private loaders
    # ------------------------------------------------------------------

    @staticmethod
    def _split_documents(documents: list[Document]) -> list[Document]:
        """Split documents into smaller chunks while preserving metadata."""
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=DocumentLoader.CHUNK_SIZE,
            chunk_overlap=DocumentLoader.CHUNK_OVERLAP,
            separators=["\n\n", "\n", ".", " ", ""],  # Split on these in order of preference
        )
        
        split_docs = []
        for doc in documents:
            # Split the document
            chunks = splitter.split_documents([doc])
            # Preserve original metadata for each chunk
            for chunk in chunks:
                chunk.metadata = doc.metadata.copy()
            split_docs.extend(chunks)
        
        return split_docs

    @staticmethod
    def _load_pdf(filepath: str) -> list[Document]:
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ImportError("pypdf is required to load PDF files. Run: pip install pypdf")

        reader = PdfReader(filepath)
        documents = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                documents.append(Document(page_content=text, metadata={"page": i + 1}))
        return documents

    @staticmethod
    def _load_docx(filepath: str) -> list[Document]:
        try:
            import docx
        except ImportError:
            raise ImportError("python-docx is required to load DOCX files. Run: pip install python-docx")

        doc = docx.Document(filepath)
        text = "\n".join(para.text for para in doc.paragraphs if para.text.strip())
        if text:
            return [Document(page_content=text)]
        return []

    @staticmethod
    def _load_csv(filepath: str) -> list[Document]:
        try:
            import pandas as pd
            df = pd.read_csv(filepath)
            documents = []
            for _, row in df.iterrows():
                text = " | ".join(
                    f"{col}: {val}" for col, val in row.items() if pd.notna(val)
                )
                if text.strip():
                    documents.append(Document(page_content=text))
            return documents
        except Exception:
            # Fallback to stdlib csv — row-wise
            with open(filepath, newline="", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                rows = list(reader)
            if not rows:
                return []
            headers = rows[0]
            documents = []
            for row in rows[1:]:
                text = " | ".join(
                    f"{h}: {v}" for h, v in zip(headers, row) if v.strip()
                )
                if text.strip():
                    documents.append(Document(page_content=text))
            return documents

    @staticmethod
    def _load_excel(filepath: str) -> list[Document]:
        try:
            import pandas as pd
        except ImportError:
            raise ImportError("pandas and openpyxl are required. Run: pip install pandas openpyxl")

        xl = pd.ExcelFile(filepath)
        documents = []
        for sheet_name in xl.sheet_names:
            df = xl.parse(sheet_name)
            for _, row in df.iterrows():
                text = f"Sheet: {sheet_name} | " + " | ".join(
                    f"{col}: {val}" for col, val in row.items() if pd.notna(val)
                )
                if text.strip():
                    documents.append(Document(page_content=text))
        return documents

    @staticmethod
    def _load_json(filepath: str) -> list[Document]:
        with open(filepath, encoding="utf-8", errors="replace") as f:
            data = json.load(f)

        documents = []
        items = data if isinstance(data, list) else [data]
        for item in items:
            if isinstance(item, dict):
                text = " | ".join(f"{k}: {v}" for k, v in item.items() if v is not None)
            else:
                text = str(item)
            if text.strip():
                documents.append(Document(page_content=text))
        return documents

    @staticmethod
    def _load_text(filepath: str) -> list[Document]:
        with open(filepath, encoding="utf-8", errors="replace") as f:
            text = f.read()
        if text.strip():
            return [Document(page_content=text)]
        return []
